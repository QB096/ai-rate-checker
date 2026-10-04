#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
小说 AI 率检测工具（LLM 判别版）
- 按章节切分 txt/md 小说
- 每章按字数分块，调用 DeepSeek 逐块判别 AI 生成概率
- 输出 Markdown + Excel 检测报告

用法:
  python ai_rate_checker.py 小说.txt
  python ai_rate_checker.py novel.md --output ./reports --threshold 70
"""

import argparse
import asyncio
import json
import re
import sys
import time
from pathlib import Path

import httpx

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    openpyxl = None

TOOL_DIR = Path(__file__).resolve().parent
CONFIG_FILE = TOOL_DIR / "config.json"

DEFAULT_CONFIG = {
    "base_url": "https://api.deepseek.com",
    "api_key": "",
    "model": "deepseek-flash",
    "concurrency": 4,
    "chunk_chars": 1200,
    "threshold": 70,
}

# ---------------------------------------------------------------- 章节切分

CHAP_PATTERNS = [
    re.compile(r"^\s*第\s*[零一二三四五六七八九十百千万两0-9０-９]+\s*[章回节卷集部篇]([　\s\.、:：].*)?$"),
    re.compile(r"^\s*Chapter\s+\d+.*$", re.IGNORECASE),
    re.compile(r"^\s*#{1,3}\s+\S.*"),
    re.compile(r"^\s*[（(]?\d{1,4}[）)]?[.、．]\s*\S.*"),
    re.compile(r"^\s*[一二三四五六七八九十]+\s*[、.]\s*\S.*"),
]


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def split_chapters(text: str):
    lines = text.splitlines()
    chapters = []
    cur_title, cur_lines = None, []
    for line in lines:
        if any(p.match(line) for p in CHAP_PATTERNS):
            if cur_lines and any(s.strip() for s in cur_lines):
                chapters.append((cur_title or "（前言/开头）", "\n".join(cur_lines)))
            cur_title = line.strip()
            cur_lines = []
        else:
            cur_lines.append(line)
    if cur_lines and any(s.strip() for s in cur_lines):
        chapters.append((cur_title or ("全文" if not chapters else "（结尾）"), "\n".join(cur_lines)))
    if not chapters:
        chapters = [("全文", text)]
    return chapters


def split_chunks(body: str, chunk_chars: int):
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n", body) if p.strip()]
    chunks, buf, size = [], [], 0
    for p in paras:
        if size + len(p) > chunk_chars and buf:
            chunks.append("\n".join(buf))
            buf, size = [], 0
        buf.append(p)
        size += len(p)
    if buf:
        chunks.append("\n".join(buf))
    return [c for c in chunks if len(c) >= 80] or ([body.strip()] if body.strip() else [])


# ---------------------------------------------------------------- LLM 判别

PROMPT = """你是一名资深的文本取证与写作风格分析专家，专门检测小说/文学类文本是否由 AI 生成。
只基于文本内在的语言特征判断，不要考虑题材、剧情合理性、内容质量。

分析维度（每项 0-100 分，越高越像 AI）：
1. 句式节奏：句长是否高度均匀、缺少人类写作长短句的随机波动
2. 结构模板：是否总分总、段落结构雷同、排比工整堆砌
3. 词汇特征：是否大量使用 AI 高频表达（仿佛、宛如、交织、弥漫、缓缓、悄然、深邃、璀璨、命运的齿轮、眼中闪过一丝等）、四字格密集工整
4. 过渡连接：是否滥用"然而""与此同时""值得注意的是""更重要的是"等书面过渡词
5. 情感细节：情感是否直给而非靠动作/细节展示，比喻是否陈词滥调
6. 人类瑕疵：真实人类写作的口语、省略、跳脱、不规则标点、即兴感是否缺失

判定基准：
0-20 几乎肯定是人类写的；21-40 偏人类；41-60 难以判断；61-80 偏向 AI；81-100 几乎肯定是 AI 写的。

只输出如下 JSON，不要输出任何其他内容：
{{"ai_score": 整数0-100, "confidence": 整数0-100, "dimensions": {{"句式节奏":0-100,"结构模板":0-100,"词汇特征":0-100,"过渡连接":0-100,"情感细节":0-100,"人类瑕疵缺失":0-100}}, "reasons": "简要判断依据(60字内)", "evidence": ["原文证据片段1", "证据片段2"]}}

待检测文本：
{__TEXT__}"""


def build_prompt(text: str) -> str:
    return PROMPT.format(__TEXT__=text)


def extract_json(s: str) -> dict:
    s = s.strip()
    s = re.sub(r"^```(?:json)?|```$", "", s, flags=re.MULTILINE).strip()
    try:
        return json.loads(s)
    except json.JSONDecodeError:
        start, end = s.find("{"), s.rfind("}")
        if start >= 0 and end > start:
            return json.loads(s[start:end + 1])
        raise


async def detect_chunk(client: httpx.AsyncClient, sem: asyncio.Semaphore, cfg: dict,
                       text: str, retries: int = 3) -> dict:
    url = cfg["base_url"].rstrip("/") + "/chat/completions"
    headers = {"Authorization": f"Bearer {cfg['api_key']}", "Content-Type": "application/json"}
    payload = {
        "model": cfg["model"],
        "messages": [{"role": "user", "content": build_prompt(text)}],
        "temperature": 0,
        "max_tokens": 4000,
    }
    async with sem:
        last_err = None
        for attempt in range(retries):
            try:
                r = await client.post(url, headers=headers, json=payload, timeout=120)
                if r.status_code in (429, 500, 502, 503, 504):
                    raise RuntimeError(f"HTTP {r.status_code}")
                r.raise_for_status()
                content = r.json()["choices"][0]["message"]["content"]
                return normalize_result(extract_json(content))
            except Exception as e:  # noqa: BLE001
                last_err = e
                await asyncio.sleep(1.5 * (2 ** attempt))
    return {"ai_score": -1, "confidence": 0, "dimensions": {}, "reasons": f"检测失败: {last_err}", "evidence": []}


def normalize_result(d: dict) -> dict:
    def clamp(v):
        try:
            return max(0, min(100, int(round(float(v)))))
        except (TypeError, ValueError):
            return 50
    return {
        "ai_score": clamp(d.get("ai_score")),
        "confidence": clamp(d.get("confidence", 50)),
        "dimensions": {k: clamp(v) for k, v in (d.get("dimensions") or {}).items()},
        "reasons": str(d.get("reasons", ""))[:300],
        "evidence": [str(x)[:120] for x in (d.get("evidence") or [])[:3]],
    }


# ---------------------------------------------------------------- 主流程

async def process_file(path: Path, cfg: dict, args) -> dict:
    text = read_text(path)
    chapters = split_chapters(text)
    if args.max_chapters:
        chapters = chapters[:args.max_chapters]
    total_chars = sum(len(re.sub(r"\s", "", b)) for _, b in chapters)

    blocks = []
    for ci, (title, body) in enumerate(chapters):
        for bi, chunk in enumerate(split_chunks(body, cfg["chunk_chars"])):
            blocks.append({"ci": ci, "bi": bi, "chapter": title, "text": chunk,
                           "chars": len(re.sub(r"\s", "", chunk))})
    if args.max_chunks:
        blocks = blocks[:args.max_chunks]

    print(f"[{path.name}] 章节 {len(chapters)} | 待检测块 {len(blocks)} | 正文 {total_chars} 字")

    sem = asyncio.Semaphore(cfg["concurrency"])
    async with httpx.AsyncClient() as client:
        tasks = [detect_chunk(client, sem, cfg, b["text"]) for b in blocks]
        done, futures = [], asyncio.as_completed(tasks)
        for i, fut in enumerate(futures, 1):
            done.append(await fut)
            if i % 5 == 0 or i == len(tasks):
                print(f"  进度 {i}/{len(tasks)}", end="\r")
    print()

    for b, r in zip(blocks, done):
        b["result"] = r

    for ci, (title, body) in enumerate(chapters):
        cb = [b for b in blocks if b["ci"] == ci]
        weight = sum(b["chars"] for b in cb) or 1
        score = sum(max(b["result"]["ai_score"], 0) * b["chars"] for b in cb) / weight
        for b in cb:
            b["chapter_score"] = score
        chapters[ci] = (title, body, score, weight)

    ok = [b for b in blocks if b["result"]["ai_score"] >= 0]
    weight = sum(b["chars"] for b in ok) or 1
    overall = sum(b["result"]["ai_score"] * b["chars"] for b in ok) / weight if ok else 0.0

    return {
        "file": path.name,
        "path": str(path),
        "total_chars": total_chars,
        "overall": round(overall, 1),
        "chapters": [{"title": t, "score": round(s, 1), "chars": w,
                      "flag": round(s, 1) >= args.threshold}
                     for t, _, s, w in chapters],
        "blocks": blocks,
    }


# ---------------------------------------------------------------- 报告

def verdict(score: float, th: int) -> str:
    if score >= 85:
        return "极高（疑似 AI 生成）"
    if score >= th:
        return "偏高（疑似 AI 润色/生成）"
    if score >= 45:
        return "中等（人机混合嫌疑）"
    if score >= 25:
        return "较低（偏人类）"
    return "低（基本为人类创作）"


def write_markdown(results, th: int, out: Path):
    lines = ["# 小说 AI 率检测报告", "",
             f"- 生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}",
             f"- 检测方式：LLM 判别（DeepSeek）+ 按字数加权",
             f"- 高风险阈值：{th} 分", ""]
    avg = sum(r["overall"] for r in results) / len(results)
    lines += ["## 总览", "",
              "| 文件 | 总 AI 率 | 正文字数 | 判定 |",
              "|---|---|---|---|"]
    for r in results:
        lines.append(f"| {r['file']} | {r['overall']}% | {r['total_chars']} | {verdict(r['overall'], th)} |")
    lines += ["", f"**综合 AI 率：{avg:.1f}%** —— {verdict(avg, th)}", ""]

    for r in results:
        lines += [f"## 《{r['file']}》章节明细", "",
                  "| 章节 | 字数 | AI 率 | 风险 |",
                  "|---|---|---|---|"]
        for c in r["chapters"]:
            flag = "🔴 高风险" if c["flag"] else ("🟡 留意" if c["score"] >= 45 else "✅ 正常")
            title = c["title"] if len(c["title"]) <= 40 else c["title"][:40] + "…"
            lines.append(f"| {title} | {c['chars']} | {c['score']}% | {flag} |")

        risky = [b for b in r["blocks"] if b["result"]["ai_score"] >= th]
        if risky:
            lines += ["", f"### 高风险段落（{len(risky)} 处，阈值 {th}）", ""]
            for b in sorted(risky, key=lambda x: -x["result"]["ai_score"])[:30]:
                res = b["result"]
                snippet = b["text"][:80].replace("\n", " ")
                lines += [f"**{b['chapter']} · 段块 #{b['bi']} · AI 率 {res['ai_score']}%（置信度 {res['confidence']}%）**",
                          f"> {snippet}…",
                          f"> 判断依据：{res['reasons']}"]
                if res["evidence"]:
                    lines.append("> 证据：" + "；".join(f"“{e}”" for e in res["evidence"]))
                lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")


def write_excel(results, th: int, out: Path):
    wb = openpyxl.Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "总览"
    sheet(ws, ["文件", "总AI率(%)", "正文字数", "章节数", "高风险章节数", "判定"],
          [[r["file"], r["overall"], r["total_chars"], len(r["chapters"]),
            sum(1 for c in r["chapters"] if c["flag"]), verdict(r["overall"], th)] for r in results],
          [36, 12, 12, 10, 14, 26])

    ws = wb.create_sheet("章节明细")
    rows = []
    for r in results:
        for c in r["chapters"]:
            rows.append([r["file"], c["title"], c["chars"], c["score"],
                         "高风险" if c["flag"] else ("留意" if c["score"] >= 45 else "正常")])
    sheet(ws, ["文件", "章节", "字数", "AI率(%)", "风险"], rows, [30, 50, 10, 10, 10])

    ws = wb.create_sheet("段落明细")
    rows = []
    for r in results:
        for b in r["blocks"]:
            res = b["result"]
            dims = res.get("dimensions") or {}
            rows.append([r["file"], b["chapter"], b["bi"], b["chars"],
                         res["ai_score"], res["confidence"],
                         dims.get("句式节奏", ""), dims.get("结构模板", ""), dims.get("词汇特征", ""),
                         dims.get("过渡连接", ""), dims.get("情感细节", ""), dims.get("人类瑕疵缺失", ""),
                         res["reasons"], "；".join(res["evidence"]),
                         b["text"][:200].replace("\n", " ")])
    sheet(ws, ["文件", "章节", "块#", "字数", "AI率", "置信度", "句式节奏", "结构模板",
               "词汇特征", "过渡连接", "情感细节", "人类瑕疵缺失", "判断依据", "证据", "片段摘录"],
          rows, [24, 34, 6, 8, 8, 8, 10, 10, 10, 10, 10, 14, 46, 46, 60])

    ws = wb.create_sheet("高风险片段")
    rows = []
    for r in results:
        for b in r["blocks"]:
            res = b["result"]
            if res["ai_score"] >= th:
                rows.append([r["file"], b["chapter"], res["ai_score"], res["confidence"],
                             res["reasons"], "；".join(res["evidence"]),
                             b["text"][:500].replace("\n", " ")])
    rows.sort(key=lambda x: -x[2])
    sheet(ws, ["文件", "章节", "AI率", "置信度", "判断依据", "证据", "原文摘录"],
          rows, [24, 34, 8, 8, 46, 46, 80])

    wb.save(out)


# ---------------------------------------------------------------- CLI

def load_config(args) -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_FILE.exists():
        try:
            cfg.update(json.loads(CONFIG_FILE.read_text(encoding="utf-8")))
        except json.JSONDecodeError:
            print(f"警告: {CONFIG_FILE} 解析失败，使用默认配置")
    import os
    cfg["api_key"] = os.environ.get("DEEPSEEK_API_KEY") or cfg["api_key"]
    if args.base_url:
        cfg["base_url"] = args.base_url
    if args.api_key:
        cfg["api_key"] = args.api_key
    if args.model:
        cfg["model"] = args.model
    if args.concurrency:
        cfg["concurrency"] = args.concurrency
    if args.chunk_chars:
        cfg["chunk_chars"] = args.chunk_chars
    if not cfg["api_key"]:
        sys.exit("错误: 未配置 API Key（config.json 的 api_key 或环境变量 DEEPSEEK_API_KEY）")
    return cfg


def main():
    ap = argparse.ArgumentParser(description="小说 AI 率检测（LLM 判别）")
    ap.add_argument("files", nargs="+", help="txt/md 小说文件")
    ap.add_argument("--output", "-o", help="报告输出目录（默认与输入同目录）")
    ap.add_argument("--threshold", "-t", type=int, default=DEFAULT_CONFIG["threshold"],
                    help="高风险阈值，默认 70")
    ap.add_argument("--chunk-chars", type=int, default=None, help="每块字数，默认 1200")
    ap.add_argument("--concurrency", "-c", type=int, default=None, help="并发数，默认 4")
    ap.add_argument("--model", default=None, help="模型名，默认 deepseek-flash")
    ap.add_argument("--base-url", default=None, help="API 地址")
    ap.add_argument("--api-key", default=None, help="API Key")
    ap.add_argument("--max-chunks", type=int, default=None, help="最多检测块数（调试用）")
    ap.add_argument("--max-chapters", type=int, default=None, help="最多检测章节数（调试用）")
    ap.add_argument("--format", choices=["both", "md", "xlsx"], default="both", help="报告格式")
    args = ap.parse_args()

    cfg = load_config(args)
    out_dir = Path(args.output) if args.output else None

    results = []
    for f in args.files:
        p = Path(f)
        if not p.exists():
            print(f"跳过（文件不存在）: {p}")
            continue
        results.append(asyncio.run(process_file(p, cfg, args)))
    if not results:
        sys.exit("没有可检测的文件")

    for r in results:
        p = Path(r["path"])
        od = out_dir or p.parent
        od.mkdir(parents=True, exist_ok=True)
        stem = p.stem
        if args.format in ("both", "md"):
            md = od / f"{stem}_AI率报告.md"
            write_markdown(results if len(results) == 1 else [r], args.threshold, md)
            print(f"报告已生成: {md}")
        if args.format in ("both", "xlsx"):
            if openpyxl is None:
                print("未安装 openpyxl，跳过 Excel 报告（pip install openpyxl）")
            else:
                xl = od / f"{stem}_AI率报告.xlsx"
                write_excel([r], args.threshold, xl)
                print(f"报告已生成: {xl}")
        print(f"  《{r['file']}》总 AI 率: {r['overall']}% —— {verdict(r['overall'], args.threshold)}")


if __name__ == "__main__":
    main()
